import io
import uuid
from datetime import date

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import or_
from sqlalchemy.orm import Session, selectinload

from app.core.acesso import requer_acesso
from app.core.database import get_db
from app.models.categoria_material import CategoriaMaterial
from app.models.execucao_reclamacao import ExecucaoReclamacao, ItemExecucaoMaterial
from app.models.lampada import PotenciaLampada, TipoLampada
from app.models.prefeitura import Prefeitura
from app.models.reclamacao import Reclamacao

router = APIRouter(prefix="/api/relatorios", tags=["relatorios"])


class LinhaPontosAtendidos(BaseModel):
    data: date
    codigo_reclamacao: str
    bairro: str
    logradouro: str
    luminarias_w: str
    lampadas_tipo: str
    condutores: str
    por_categoria: dict[str, float]
    pontos: float


def _com_dados(query):
    return query.options(
        selectinload(ExecucaoReclamacao.itens).selectinload(ItemExecucaoMaterial.material),
        selectinload(ExecucaoReclamacao.reclamacao).selectinload(Reclamacao.bairro),
    )


@router.get(
    "/pontos-atendidos",
    response_model=list[LinhaPontosAtendidos],
    dependencies=[Depends(requer_acesso("relatorios", "use"))],
)
def pontos_atendidos(
    prefeitura_id: uuid.UUID,
    data_inicio: date,
    data_fim: date,
    db: Session = Depends(get_db),
):
    """Uma linha por execução de reclamação da prefeitura no período,
    agregando a quantidade instalada de materiais por categoria (uma coluna
    por categoria do catálogo) e as lâmpadas instaladas (quantidade-potência)."""
    categorias_codigos = [c for (c,) in db.query(CategoriaMaterial.codigo).all()]

    execucoes = (
        _com_dados(
            db.query(ExecucaoReclamacao)
            .join(Reclamacao, ExecucaoReclamacao.reclamacao_id == Reclamacao.id)
            .filter(
                Reclamacao.prefeitura_id == prefeitura_id,
                ExecucaoReclamacao.data_execucao >= data_inicio,
                ExecucaoReclamacao.data_execucao <= data_fim,
            )
        )
        .order_by(ExecucaoReclamacao.data_execucao, Reclamacao.bairro_id)
        .all()
    )

    potencias = {p.id: p.valor_w for p in db.query(PotenciaLampada).all()}
    tipos_lampada = {t.id: t.nome for t in db.query(TipoLampada).all()}

    linhas = []
    for ex in execucoes:
        somas = {codigo: 0.0 for codigo in categorias_codigos}
        qde_por_potencia_lampada = {}
        qde_por_tipo_lampada = {}
        qde_por_descricao_condutor = {}
        for item in ex.itens:
            categoria = item.material.categoria if item.material else None
            if categoria in somas:
                somas[categoria] += float(item.quantidade_instalada or 0)
            if categoria == "LAMPADA" and item.quantidade_instalada:
                potencia = potencias.get(item.potencia_lampada_id)
                sufixo = f"{float(potencia):g}W" if potencia else "—"
                qde_por_potencia_lampada[sufixo] = (
                    qde_por_potencia_lampada.get(sufixo, 0.0) + float(item.quantidade_instalada)
                )
                tipo = tipos_lampada.get(item.tipo_lampada_id) or "—"
                qde_por_tipo_lampada[tipo] = qde_por_tipo_lampada.get(tipo, 0.0) + float(item.quantidade_instalada)
            if categoria == "CONDUTOR" and item.quantidade_instalada:
                descricao = item.material.nome if item.material else "—"
                qde_por_descricao_condutor[descricao] = (
                    qde_por_descricao_condutor.get(descricao, 0.0) + float(item.quantidade_instalada)
                )

        # Mesma potência, mesmo tipo de lâmpada ou mesma descrição (condutor)
        # soma a quantidade em vez de listar um item por linha lançada.
        lampadas_potencia = [f"{qde:g}-{pot}" for pot, qde in qde_por_potencia_lampada.items()]
        lampadas_tipo = [f"{qde:g}-{tipo}" for tipo, qde in qde_por_tipo_lampada.items()]
        condutores = [f"{qde:g}-{descricao}" for descricao, qde in qde_por_descricao_condutor.items()]

        reclamacao = ex.reclamacao
        logradouro = reclamacao.logradouro or "—"
        if reclamacao.numero:
            logradouro = f"{logradouro}, Nº {reclamacao.numero}"

        linhas.append(
            LinhaPontosAtendidos(
                data=ex.data_execucao,
                codigo_reclamacao=reclamacao.codigo,
                bairro=reclamacao.bairro.nome if reclamacao.bairro else "—",
                logradouro=logradouro,
                luminarias_w=", ".join(lampadas_potencia),
                lampadas_tipo=", ".join(lampadas_tipo),
                condutores=", ".join(condutores),
                por_categoria=somas,
                pontos=ex.pontos,
            )
        )
    return linhas


class LinhaMaterialGasto(BaseModel):
    codigo: str
    material: str
    categoria: str
    unidade: str
    quantidade: float  # instalada
    quantidade_retirada: float = 0
    quantidade_substituida: float = 0


@router.get(
    "/materiais-gastos",
    response_model=list[LinhaMaterialGasto],
    dependencies=[Depends(requer_acesso("relatorios", "use"))],
)
def materiais_gastos(
    prefeitura_id: uuid.UUID,
    data_inicio: date,
    data_fim: date,
    db: Session = Depends(get_db),
):
    return _somar_materiais(prefeitura_id, data_inicio, data_fim, db)


def _somar_materiais(prefeitura_id: uuid.UUID, data_inicio: date, data_fim: date, db: Session) -> list[LinhaMaterialGasto]:
    """Materiais lançados nas execuções da prefeitura no período, com as
    quantidades instalada, retirada e substituída somadas por material (sem
    quebrar por reclamação). Lâmpadas ficam separadas por tipo e potência,
    pois são itens diferentes no estoque."""
    itens = (
        db.query(ItemExecucaoMaterial)
        .join(ExecucaoReclamacao, ExecucaoReclamacao.id == ItemExecucaoMaterial.execucao_id)
        .join(Reclamacao, Reclamacao.id == ExecucaoReclamacao.reclamacao_id)
        .options(selectinload(ItemExecucaoMaterial.material))
        .filter(
            Reclamacao.prefeitura_id == prefeitura_id,
            ExecucaoReclamacao.data_execucao >= data_inicio,
            ExecucaoReclamacao.data_execucao <= data_fim,
            or_(
                ItemExecucaoMaterial.quantidade_instalada > 0,
                ItemExecucaoMaterial.quantidade_retirada > 0,
                ItemExecucaoMaterial.quantidade_substituida > 0,
            ),
        )
        .all()
    )

    categorias = {c.codigo: c.nome for c in db.query(CategoriaMaterial).all()}
    potencias = {p.id: p.valor_w for p in db.query(PotenciaLampada).all()}
    tipos_lampada = {t.id: t.nome for t in db.query(TipoLampada).all()}

    somas: dict[tuple, list[float]] = {}
    for item in itens:
        material = item.material
        if not material:
            continue
        descricao = material.nome
        if material.categoria == "LAMPADA":
            tipo = tipos_lampada.get(item.tipo_lampada_id)
            potencia = potencias.get(item.potencia_lampada_id)
            detalhe = " ".join(p for p in (tipo, f"{float(potencia):g}W" if potencia else None) if p)
            if detalhe:
                descricao = f"{material.nome} — {detalhe}"
        chave = (categorias.get(material.categoria, material.categoria), descricao, material.codigo, material.unidade)
        total = somas.setdefault(chave, [0.0, 0.0, 0.0])
        total[0] += float(item.quantidade_instalada or 0)
        total[1] += float(item.quantidade_retirada or 0)
        total[2] += float(item.quantidade_substituida or 0)

    return [
        LinhaMaterialGasto(
            codigo=codigo,
            material=descricao,
            categoria=categoria,
            unidade=unidade,
            quantidade=inst,
            quantidade_retirada=ret,
            quantidade_substituida=subst,
        )
        for (categoria, descricao, codigo, unidade), (inst, ret, subst) in sorted(
            somas.items(), key=lambda kv: (kv[0][0], kv[0][1])
        )
    ]


@router.get(
    "/materiais-gastos/excel",
    dependencies=[Depends(requer_acesso("relatorios", "use"))],
)
def materiais_gastos_excel(
    prefeitura_id: uuid.UUID,
    data_inicio: date,
    data_fim: date,
    db: Session = Depends(get_db),
):
    """Mesmo relatório de materiais gastos, em planilha .xlsx."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    prefeitura = db.get(Prefeitura, prefeitura_id)
    linhas = _somar_materiais(prefeitura_id, data_inicio, data_fim, db)

    wb = Workbook()
    ws = wb.active
    ws.title = "Materiais gastos"
    ws["A1"] = "RELATÓRIO DE MATERIAIS GASTOS"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = prefeitura.nome if prefeitura else ""
    ws["A3"] = f"Período: {data_inicio:%d/%m/%Y} a {data_fim:%d/%m/%Y}"

    cabecalho = ["Código", "Material", "Categoria", "Unid.", "Qtd. instalada", "Qtd. retirada", "Qtd. substituída"]
    for col, titulo in enumerate(cabecalho, start=1):
        c = ws.cell(row=5, column=col, value=titulo)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E5EBF7")
        c.alignment = Alignment(horizontal="center" if col >= 4 else "left")
    for i, l in enumerate(linhas, start=6):
        ws.cell(row=i, column=1, value=l.codigo)
        ws.cell(row=i, column=2, value=l.material)
        ws.cell(row=i, column=3, value=l.categoria)
        ws.cell(row=i, column=4, value=l.unidade).alignment = Alignment(horizontal="center")
        # Números de verdade (somáveis no Excel). Formato sem o ".##" — ele
        # deixava uma vírgula sobrando ("1,") no Excel em português.
        for col, valor in enumerate((l.quantidade, l.quantidade_retirada, l.quantidade_substituida), start=5):
            inteiro = float(valor).is_integer()
            qtd = ws.cell(row=i, column=col, value=int(valor) if inteiro else valor)
            qtd.number_format = "#,##0" if inteiro else "#,##0.00"
            qtd.alignment = Alignment(horizontal="center")
    for col, largura in enumerate([16, 46, 18, 8, 16, 16, 18], start=1):
        ws.column_dimensions[get_column_letter(col)].width = largura

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="materiais_gastos.xlsx"'},
    )


# Colunas do relatório — espelham frontend/src/relatorioColunas.js (mantenha os dois iguais).
_ABREVIACOES = {"CONECTOR": "Conx", "ISOLADOR": "Isol", "CONDUTOR": "Cond", "LAMPADA": "Lamp"}
_MESCLADAS_EM_OUTROS = ["ISOLANTES", "BRACO", "FERRAGENS", "ISOLADOR", "POSTE"]
_ORDEM_CATEGORIAS = ["RELE", "BASE", "CONDUTOR", "CONECTOR", "LUMINARIA", "OUTROS", "REFLETOR"]


@router.get(
    "/pontos-atendidos/excel",
    dependencies=[Depends(requer_acesso("relatorios", "use"))],
)
def pontos_atendidos_excel(
    prefeitura_id: uuid.UUID,
    data_inicio: date,
    data_fim: date,
    db: Session = Depends(get_db),
):
    """Relatório de pontos atendidos em planilha .xlsx, com as mesmas colunas
    da tela (Lamp, Pot.(W), uma coluna por categoria, Pontos) e linha de totais."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    prefeitura = db.get(Prefeitura, prefeitura_id)
    linhas = pontos_atendidos(prefeitura_id, data_inicio, data_fim, db)

    categorias = db.query(CategoriaMaterial).all()
    tem_lampada = any(c.codigo == "LAMPADA" for c in categorias)
    visiveis = sorted(
        (c for c in categorias if c.codigo not in _MESCLADAS_EM_OUTROS and c.codigo != "LAMPADA"),
        key=lambda c: _ORDEM_CATEGORIAS.index(c.codigo) if c.codigo in _ORDEM_CATEGORIAS else -1,
    )

    def valor_categoria(linha: LinhaPontosAtendidos, codigo: str) -> float:
        base = float(linha.por_categoria.get(codigo, 0) or 0)
        if codigo == "OUTROS":
            base += sum(float(linha.por_categoria.get(m, 0) or 0) for m in _MESCLADAS_EM_OUTROS)
        return base

    def formato(v: float) -> str:
        # "—" no lugar de zero; sem ".##" (deixava vírgula sobrando no Excel em português).
        return '#,##0;-#,##0;"—"' if float(v).is_integer() else '#,##0.00;-#,##0.00;"—"'

    wb = Workbook()
    ws = wb.active
    ws.title = "Pontos atendidos"
    ws["A1"] = "RELATÓRIO DE PONTOS ATENDIDOS"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = prefeitura.nome if prefeitura else ""
    ws["A3"] = f"Período: {data_inicio:%d/%m/%Y} a {data_fim:%d/%m/%Y}"

    cabecalho = ["Código", "Data", "Bairro", "Logradouro"]
    if tem_lampada:
        cabecalho.append("Lamp")
    cabecalho.append("Pot.(W)")
    cabecalho += [_ABREVIACOES.get(c.codigo, c.nome) for c in visiveis]
    cabecalho.append("Pontos")
    for col, titulo in enumerate(cabecalho, start=1):
        c = ws.cell(row=5, column=col, value=titulo)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E5EBF7")
        c.alignment = Alignment(horizontal="left" if col <= 4 else "center")

    centro = Alignment(horizontal="center", wrap_text=True)
    totais = {c.codigo: 0.0 for c in visiveis}
    total_pontos = 0.0
    linha_excel = 6
    for l in linhas:
        valores = [l.codigo_reclamacao, l.data, l.bairro, l.logradouro]
        if tem_lampada:
            valores.append(l.lampadas_tipo or "—")
        valores.append(l.luminarias_w or "—")
        for col, v in enumerate(valores, start=1):
            c = ws.cell(row=linha_excel, column=col, value=v)
            if col == 2:
                c.number_format = "DD/MM/YYYY"
                c.alignment = Alignment(horizontal="center")
            elif col > 4:
                c.alignment = centro
        col = len(valores) + 1
        for cat in visiveis:
            if cat.codigo == "CONDUTOR":
                c = ws.cell(row=linha_excel, column=col, value=l.condutores or "—")
                c.alignment = centro
            else:
                v = valor_categoria(l, cat.codigo)
                totais[cat.codigo] += v
                c = ws.cell(row=linha_excel, column=col, value=int(v) if float(v).is_integer() else v)
                c.number_format = formato(v)
                c.alignment = centro
            col += 1
        pontos = float(l.pontos or 0)
        total_pontos += pontos
        c = ws.cell(row=linha_excel, column=col, value=int(pontos) if pontos.is_integer() else pontos)
        c.number_format = formato(pontos)
        c.alignment = centro
        linha_excel += 1

    # Totais (Cond é texto, então fica "—", como na tela).
    ws.cell(row=linha_excel, column=1, value="Totais").font = Font(bold=True)
    col = 6 + (1 if tem_lampada else 0)  # depois de Pot.(W)
    for cat in visiveis:
        if cat.codigo == "CONDUTOR":
            v = "—"
            c = ws.cell(row=linha_excel, column=col, value=v)
        else:
            t = totais[cat.codigo]
            c = ws.cell(row=linha_excel, column=col, value=int(t) if t.is_integer() else t)
            c.number_format = formato(t)
        c.font = Font(bold=True)
        c.alignment = centro
        col += 1
    c = ws.cell(row=linha_excel, column=col, value=int(total_pontos) if total_pontos.is_integer() else total_pontos)
    c.number_format = formato(total_pontos)
    c.font = Font(bold=True)
    c.alignment = centro
    for k in range(1, col + 1):
        ws.cell(row=linha_excel, column=k).fill = PatternFill("solid", fgColor="EEF2FA")

    larguras = [20, 12, 22, 36] + ([14] if tem_lampada else []) + [14] + [11] * len(visiveis) + [10]
    for k, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(k)].width = largura

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="pontos_atendidos.xlsx"'},
    )
