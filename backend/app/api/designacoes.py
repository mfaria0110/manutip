import io
import uuid
from datetime import date

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session, selectinload

from app.core.acesso import requer_acesso
from app.core.database import get_db
from app.models.designacao import DesignacaoReclamacao
from app.models.pessoal import EquipeDia, EquipeMembro
from app.models.prefeitura import Prefeitura
from app.models.reclamacao import Reclamacao

router = APIRouter(prefix="/api/designacoes", tags=["designacoes"])

# Reclamações que ainda podem ser designadas a uma equipe.
STATUS_DESIGNAVEIS = ("ABERTA", "EM_ANDAMENTO")


class DesignacaoOut(BaseModel):
    id: uuid.UUID
    equipe_dia_id: uuid.UUID
    reclamacao_id: uuid.UUID
    data: date

    class Config:
        from_attributes = True


class DesignacaoItem(BaseModel):
    reclamacao_id: uuid.UUID
    equipe_dia_id: uuid.UUID


class DesignacaoSalvar(BaseModel):
    # Designação é por prefeitura e dia: o formulário enxerga as reclamações
    # abertas de uma prefeitura, então salvar substitui só as dessa
    # prefeitura naquele dia (a data é a do roteiro, independe de quando a
    # equipe foi cadastrada).
    prefeitura_id: uuid.UUID
    data: date
    itens: list[DesignacaoItem] = []


@router.get("", response_model=list[DesignacaoOut], dependencies=[Depends(requer_acesso("reclamacoes", "use"))])
def listar(
    data: date | None = Query(default=None),
    equipe_dia_id: uuid.UUID | None = Query(default=None),
    prefeitura_id: uuid.UUID | None = Query(default=None),
    db: Session = Depends(get_db),
):
    query = db.query(DesignacaoReclamacao)
    if data:
        query = query.filter(DesignacaoReclamacao.data == data)
    if equipe_dia_id:
        query = query.filter(DesignacaoReclamacao.equipe_dia_id == equipe_dia_id)
    if prefeitura_id:
        query = query.join(Reclamacao, Reclamacao.id == DesignacaoReclamacao.reclamacao_id).filter(
            Reclamacao.prefeitura_id == prefeitura_id
        )
    return query.all()


@router.put(
    "/dia",
    response_model=list[DesignacaoOut],
    dependencies=[Depends(requer_acesso("reclamacoes", "edit"))],
)
def salvar_do_dia(req: DesignacaoSalvar, db: Session = Depends(get_db)):
    """Grava o roteiro do dia da prefeitura: cada item liga uma reclamação
    aberta a uma equipe. Reclamação fora da lista fica sem equipe."""
    ids = [i.reclamacao_id for i in req.itens]
    if len(set(ids)) != len(ids):
        raise HTTPException(status_code=400, detail="Uma reclamação não pode ter mais de uma equipe no mesmo dia.")

    reclamacoes = {r.id: r for r in db.query(Reclamacao).filter(Reclamacao.id.in_(ids))} if ids else {}
    if len(reclamacoes) != len(ids):
        raise HTTPException(status_code=404, detail="Alguma reclamação selecionada não foi encontrada.")
    for r in reclamacoes.values():
        if r.prefeitura_id != req.prefeitura_id:
            raise HTTPException(status_code=400, detail=f"A reclamação {r.codigo} não é da prefeitura selecionada.")
        if r.status not in STATUS_DESIGNAVEIS:
            raise HTTPException(status_code=400, detail=f"A reclamação {r.codigo} não está aberta.")

    equipe_ids = {i.equipe_dia_id for i in req.itens}
    if equipe_ids and db.query(EquipeDia).filter(EquipeDia.id.in_(equipe_ids)).count() != len(equipe_ids):
        raise HTTPException(status_code=404, detail="Alguma equipe selecionada não foi encontrada.")

    # Só apaga as designações das reclamações ainda abertas: as de
    # reclamações já concluídas ficam como histórico do que foi atendido.
    abertas_da_prefeitura = db.query(Reclamacao.id).filter(
        Reclamacao.prefeitura_id == req.prefeitura_id, Reclamacao.status.in_(STATUS_DESIGNAVEIS)
    )
    db.query(DesignacaoReclamacao).filter(
        DesignacaoReclamacao.data == req.data,
        DesignacaoReclamacao.reclamacao_id.in_(abertas_da_prefeitura),
    ).delete(synchronize_session=False)
    for item in req.itens:
        db.add(
            DesignacaoReclamacao(equipe_dia_id=item.equipe_dia_id, reclamacao_id=item.reclamacao_id, data=req.data)
        )
    db.commit()
    return (
        db.query(DesignacaoReclamacao)
        .join(Reclamacao, Reclamacao.id == DesignacaoReclamacao.reclamacao_id)
        .filter(DesignacaoReclamacao.data == req.data, Reclamacao.prefeitura_id == req.prefeitura_id)
        .all()
    )


class ExcelRoteiro(BaseModel):
    prefeitura_id: uuid.UUID
    data: date
    equipe_dia_id: uuid.UUID
    reclamacao_ids: list[uuid.UUID] = []


@router.post("/excel", dependencies=[Depends(requer_acesso("reclamacoes", "use"))])
def exportar_excel(req: ExcelRoteiro, db: Session = Depends(get_db)):
    """Roteiro da equipe em planilha .xlsx, no mesmo modelo do Imprimir:
    cidade e título no topo, quadro com equipe/viatura/componentes e a tabela
    Código, Data_Rec, Bairro, Logradouro, Lamp, Pot.(W) e as categorias de
    material (Lamp até Refletor em branco e com borda, pra preencher em campo)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    from app.api.relatorios import _ABREVIACOES, _MESCLADAS_EM_OUTROS, _ORDEM_CATEGORIAS
    from app.models.categoria_material import CategoriaMaterial
    from app.models.localidade import Cidade

    prefeitura = db.get(Prefeitura, req.prefeitura_id)
    cidade = db.get(Cidade, prefeitura.cidade_id) if prefeitura and prefeitura.cidade_id else None
    equipe = (
        db.query(EquipeDia)
        .options(selectinload(EquipeDia.membros).selectinload(EquipeMembro.funcionario))
        .filter(EquipeDia.id == req.equipe_dia_id)
        .first()
    )
    if not equipe:
        raise HTTPException(status_code=404, detail="Equipe não encontrada.")
    ordem = {rid: n for n, rid in enumerate(req.reclamacao_ids)}
    reclamacoes = (
        sorted(db.query(Reclamacao).filter(Reclamacao.id.in_(req.reclamacao_ids)).all(), key=lambda r: ordem[r.id])
        if req.reclamacao_ids
        else []
    )

    categorias = db.query(CategoriaMaterial).all()
    tem_lampada = any(c.codigo == "LAMPADA" for c in categorias)
    visiveis = sorted(
        (c for c in categorias if c.codigo not in _MESCLADAS_EM_OUTROS and c.codigo != "LAMPADA"),
        key=lambda c: _ORDEM_CATEGORIAS.index(c.codigo) if c.codigo in _ORDEM_CATEGORIAS else -1,
    )

    fino = Side(style="thin", color="444444")
    borda = Border(left=fino, right=fino, top=fino, bottom=fino)
    wb = Workbook()
    ws = wb.active
    ws.title = "Roteiro"
    ws["A1"] = cidade.nome if cidade else (prefeitura.nome if prefeitura else "")
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"ROTEIRO DO DIA: {req.data:%d/%m/%Y}"
    ws["A2"].font = Font(bold=True, size=12)

    veiculo = f"{equipe.veiculo.placa} {equipe.veiculo.modelo}" if equipe.veiculo else "sem veículo"
    membros = ", ".join(m.funcionario.nome for m in equipe.membros if m.funcionario) or "sem membros"
    for linha, (rotulo, valor) in enumerate(
        (("Equipe:", equipe.nome or "Equipe"), ("Viatura:", veiculo), ("Componentes:", membros)), start=4
    ):
        ws.cell(row=linha, column=1, value=rotulo).font = Font(bold=True)
        ws.cell(row=linha, column=2, value=valor)

    cabecalho = ["Código", "Data_Rec", "Bairro", "Logradouro"]
    if tem_lampada:
        cabecalho.append("Lamp")
    cabecalho.append("Pot.(W)")
    cabecalho += [_ABREVIACOES.get(c.codigo, c.nome) for c in visiveis]
    linha_cab = 8
    for col, titulo in enumerate(cabecalho, start=1):
        c = ws.cell(row=linha_cab, column=col, value=titulo)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E5EBF7")
        c.alignment = Alignment(horizontal="left" if col <= 4 else "center")
        if col >= 5:
            c.border = borda
    linha = linha_cab + 1
    for r in reclamacoes:
        logradouro = ", ".join(p for p in (r.logradouro or "—", f"Nº {r.numero}" if r.numero else "") if p)
        valores = [r.codigo, r.data_reclamacao, r.bairro.nome if r.bairro else "—", logradouro]
        for col, v in enumerate(valores, start=1):
            c = ws.cell(row=linha, column=col, value=v)
            if col == 2:
                c.number_format = "DD/MM/YYYY"
                c.alignment = Alignment(horizontal="center")
        for col in range(5, len(cabecalho) + 1):
            ws.cell(row=linha, column=col).border = borda
        ws.row_dimensions[linha].height = 22
        linha += 1
    larguras = [20, 12, 24, 38] + ([12] if tem_lampada else []) + [12] + [11] * len(visiveis)
    for col, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(col)].width = largura

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="roteiro.xlsx"'},
    )
